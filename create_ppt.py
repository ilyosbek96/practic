from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Colors
BG_DARK = RGBColor(18, 18, 35)
ACCENT_BLUE = RGBColor(64, 158, 255)
ACCENT_GREEN = RGBColor(80, 200, 120)
ACCENT_YELLOW = RGBColor(255, 214, 0)
ACCENT_PURPLE = RGBColor(180, 100, 255)
ACCENT_ORANGE = RGBColor(255, 140, 60)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(200, 210, 230)
CODE_BG = RGBColor(28, 32, 55)
CODE_TEXT = RGBColor(130, 220, 140)


def set_bg(slide, color=BG_DARK):
    from pptx.util import Pt
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color, transparency=0):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text_box(slide, text, left, top, width, height, font_size=18, bold=False,
                  color=WHITE, align=PP_ALIGN.LEFT, font_name="Consolas"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return txBox


def add_code_block(slide, code, left, top, width, height):
    add_rect(slide, left, top, width, height, CODE_BG)
    add_text_box(slide, code, left + 0.15, top + 0.1, width - 0.3, height - 0.2,
                 font_size=13, color=CODE_TEXT, font_name="Consolas")


def add_title_slide():
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_bg(slide)

    # top accent line
    add_rect(slide, 0, 0, 13.33, 0.08, ACCENT_BLUE)
    add_rect(slide, 0, 7.42, 13.33, 0.08, ACCENT_BLUE)

    # decorative circles
    for i, (cx, cy, cr, col) in enumerate([
        (1.2, 1.5, 1.8, RGBColor(30, 50, 100)),
        (11.8, 5.8, 1.5, RGBColor(30, 80, 60)),
        (12.5, 0.8, 1.0, RGBColor(40, 30, 80)),
    ]):
        sh = slide.shapes.add_shape(9, Inches(cx - cr/2), Inches(cy - cr/2),
                                    Inches(cr), Inches(cr))
        sh.fill.solid()
        sh.fill.fore_color.rgb = col
        sh.line.fill.background()

    # Python logo placeholder text
    add_text_box(slide, "🐍", 6.2, 1.0, 1.0, 1.0, font_size=40, align=PP_ALIGN.CENTER)

    add_text_box(slide, "PYTHON", 2.5, 1.8, 8.33, 1.0, font_size=18, bold=True,
                 color=ACCENT_BLUE, align=PP_ALIGN.CENTER, font_name="Arial")
    add_text_box(slide, "LIST TO'LIQ QO'LLANMA", 1.5, 2.5, 10.33, 1.2, font_size=40,
                 bold=True, color=WHITE, align=PP_ALIGN.CENTER, font_name="Arial")

    add_rect(slide, 4.5, 3.75, 4.33, 0.05, ACCENT_BLUE)

    topics = ["Working with Lists", "List Methods", "Lambda Functions", "Enumerate, Map & Filter"]
    for i, t in enumerate(topics):
        add_text_box(slide, f"{'●'}  {t}", 3.5, 4.2 + i * 0.5, 6.33, 0.45,
                     font_size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER, font_name="Arial")

    add_text_box(slide, "Python Lists Masterclass", 4.0, 6.8, 5.33, 0.4,
                 font_size=13, color=RGBColor(100, 120, 160), align=PP_ALIGN.CENTER, font_name="Arial")


def add_section_slide(number, title, subtitle, accent_color):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_bg(slide)

    add_rect(slide, 0, 0, 13.33, 0.08, accent_color)
    add_rect(slide, 0, 7.42, 13.33, 0.08, accent_color)

    # big number bg
    add_rect(slide, 0.3, 1.5, 3.5, 4.5, RGBColor(25, 28, 50))
    add_text_box(slide, number, 0.3, 1.8, 3.5, 3.0, font_size=100, bold=True,
                 color=accent_color, align=PP_ALIGN.CENTER, font_name="Arial")

    add_text_box(slide, title, 4.3, 2.5, 8.5, 1.2, font_size=38, bold=True,
                 color=WHITE, font_name="Arial")
    add_rect(slide, 4.3, 3.8, 5.0, 0.06, accent_color)
    add_text_box(slide, subtitle, 4.3, 4.0, 8.5, 1.5, font_size=18,
                 color=LIGHT_GRAY, font_name="Arial")


def add_content_slide(title, bullets, code=None, accent_color=ACCENT_BLUE, code_height=2.5):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_bg(slide)

    add_rect(slide, 0, 0, 13.33, 0.08, accent_color)
    add_rect(slide, 0, 0.08, 13.33, 0.9, RGBColor(22, 26, 48))
    add_text_box(slide, title, 0.3, 0.1, 12.5, 0.85, font_size=24, bold=True,
                 color=WHITE, font_name="Arial")

    content_top = 1.15
    if code:
        # split screen
        # left: bullets
        y = content_top
        for bul in bullets:
            icon = "▸ " if not bul.startswith("  ") else ""
            col = LIGHT_GRAY if bul.startswith("  ") else WHITE
            fs = 15 if bul.startswith("  ") else 17
            add_text_box(slide, icon + bul.lstrip(), 0.3, y, 6.0, 0.45,
                         font_size=fs, color=col, font_name="Arial")
            y += 0.48
        # right: code
        add_text_box(slide, "Kod namunasi:", 6.7, content_top - 0.05, 6.0, 0.35,
                     font_size=14, color=ACCENT_YELLOW, bold=True, font_name="Arial")
        add_code_block(slide, code, 6.7, content_top + 0.35, 6.3, code_height)
    else:
        y = content_top
        for bul in bullets:
            if bul.startswith("##"):
                add_text_box(slide, bul[2:].strip(), 0.3, y, 12.5, 0.45,
                             font_size=18, bold=True, color=accent_color, font_name="Arial")
                y += 0.5
            elif bul.startswith("  "):
                add_text_box(slide, "  " + bul.strip(), 0.8, y, 12.0, 0.4,
                             font_size=15, color=LIGHT_GRAY, font_name="Arial")
                y += 0.44
            else:
                add_text_box(slide, "▸  " + bul, 0.3, y, 12.5, 0.44,
                             font_size=17, color=WHITE, font_name="Arial")
                y += 0.5


def add_table_slide(title, headers, rows, accent_color=ACCENT_BLUE):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_bg(slide)

    add_rect(slide, 0, 0, 13.33, 0.08, accent_color)
    add_rect(slide, 0, 0.08, 13.33, 0.9, RGBColor(22, 26, 48))
    add_text_box(slide, title, 0.3, 0.1, 12.5, 0.85, font_size=24, bold=True,
                 color=WHITE, font_name="Arial")

    col_widths = [3.0, 4.5, 5.5]
    col_starts = [0.3, 3.4, 8.0]
    row_h = 0.52
    top = 1.15

    # header
    for i, (h, w, s) in enumerate(zip(headers, col_widths, col_starts)):
        add_rect(slide, s, top, w - 0.05, row_h, accent_color)
        add_text_box(slide, h, s + 0.1, top + 0.05, w - 0.2, row_h - 0.1,
                     font_size=15, bold=True, color=WHITE, font_name="Arial")

    for r_idx, row in enumerate(rows):
        y = top + row_h + r_idx * row_h
        bg = RGBColor(25, 30, 55) if r_idx % 2 == 0 else RGBColor(30, 36, 65)
        for i, (cell, w, s) in enumerate(zip(row, col_widths, col_starts)):
            add_rect(slide, s, y, w - 0.05, row_h - 0.03, bg)
            col = CODE_TEXT if i == 0 else LIGHT_GRAY
            add_text_box(slide, cell, s + 0.1, y + 0.05, w - 0.2, row_h - 0.1,
                         font_size=13, color=col, font_name="Consolas" if i == 0 else "Arial")


# ─── SLIDES ───────────────────────────────────────────────────────────────────

# Slide 1 - Title
add_title_slide()

# Slide 2 - Section 1
add_section_slide("01", "Working with Lists",
                  "List yaratishning 2 usuli:\nLiteral va Constructor", ACCENT_BLUE)

# Slide 3 - Literal vs Constructor
add_content_slide(
    "01 › List Yaratish: Literal va Constructor",
    [
        "LITERAL usul — [] qavs ichida to'g'ridan to'g'ri",
        "  groups = [\"MIT\", \"FLEX\", \"DEVEX\"]",
        "  people = (\"Andrew\", \"John\")   # tuple ham shunday",
        "CONSTRUCTOR usul — list() funksiyasi orqali",
        "  result = list(\"Hello\")  =>  ['H','e','l','l','o']",
        "for loop bilan listni ko'rib chiqish",
        "  for team in groups:  print(team)",
    ],
    code='# Literal\ngroups = ["MIT", "FLEX", "DEVEX", "MG"]\nfor team in groups:\n    print(f"team: {team}")\n\n# Constructor\nresult = list("Hello World!")\nprint(result)\n# => [\'H\',\'e\',\'l\',\'l\',\'o\',...]\nprint(len(result))  # => 12',
    accent_color=ACCENT_BLUE
)

# Slide 4 - Slicing
add_content_slide(
    "01 › List Slicing (Kesish)",
    [
        "fruits = [\"apple\", \"orange\", \"lemon\", \"kiwi\"]",
        "fruits[0]       →  \"apple\"  (birinchi element)",
        "fruits[0:2]     →  [\"apple\", \"orange\"]  (2 kirmaydi!)",
        "fruits[::3]     →  [\"apple\", \"kiwi\"]  (har 3 qadam)",
        "fruits[::-1]    →  [\"kiwi\", \"lemon\", \"orange\", \"apple\"]",
        "Qoida:  list[start : stop : step]",
        "  start — boshlang'ich indeks (default: 0)",
        "  stop  — oxirgi indeks (kirmaydi!)",
        "  step  — qadam (default: 1, manfiy = teskari)",
    ],
    code='fruits = ["apple","orange","lemon","kiwi"]\n\na = fruits[0]      # "apple"\nb = fruits[0:2]    # ["apple","orange"]\nc = fruits[::3]    # ["apple","kiwi"]\nd = fruits[::-1]   # ["kiwi","lemon",...]\n\nprint("a:", a)\nprint("b:", b)\nprint("c:", c)\nprint("d:", d)',
    accent_color=ACCENT_BLUE,
    code_height=2.8
)

# Slide 5 - Section 2
add_section_slide("02", "List Methods",
                  "Mutable vs Immutable metodlar:\nappend, insert, pop, remove, sort...", ACCENT_GREEN)

# Slide 6 - Mutable methods table
add_table_slide(
    "02 › Mutable Metodlar (listni o'zgartiradi)",
    ["Metod", "Sintaksis", "Natija / Izoh"],
    [
        ["append()", 'letters.append("c")', "Oxiriga element qo'shadi"],
        ["insert()", 'letters.insert(0, "z")', "Belgilangan indeksga qo'shadi"],
        ["pop()", "letters.pop(size)", "O'chirib, qiymatini qaytaradi"],
        ["remove()", 'animals.remove("lion")', "Qiymat bo'yicha o'chiradi"],
        ["del", "del animals[2:4]", "Diapazon bo'yicha o'chiradi"],
        ["clear()", "animals.clear()", "Barcha elementlarni o'chiradi"],
        ["sort()", "numbers.sort()", "Kichikdan kattaga saralaydi"],
        ["sort(rev)", "numbers.sort(reverse=True)", "Kattadan kichikka saralaydi"],
    ],
    accent_color=ACCENT_GREEN
)

# Slide 7 - Immutable methods
add_content_slide(
    "02 › Immutable Metodlar (asl listni o'zgartirmaydi)",
    [
        "index() — elementning indeksini topadi",
        "  exist = animals.index(\"cat\")   # => 1",
        "sorted() — yangi saralangan list qaytaradi",
        "  new_numbs = sorted(numbs)  # numbs o'zgarmaydi!",
        "in operatori — element mavjudligini tekshiradi",
        "  if \"cat\" in animals:   # True yoki False",
    ],
    code='numbs = [2, 20, 12, 100]\n\n# Immutable - asl list saqlanadi\nnew_numbs = sorted(numbs)\nprint(numbs)     # [2, 20, 12, 100]\nprint(new_numbs) # [2, 12, 20, 100]\n\n# index() - indeks topish\nanimals = ["dog", "cat", "fish"]\nexist = animals.index("cat")\nprint(exist)  # => 1\n\n# in - tekshirish\nif "cat" in animals:\n    print("topildi!")',
    accent_color=ACCENT_GREEN,
    code_height=3.0
)

# Slide 8 - Section 3
add_section_slide("03", "Lambda Function",
                  "Kichik, anonim (nomsiz) funksiyalar\nkey=lambda pattern bilan ishlash", ACCENT_YELLOW)

# Slide 9 - Lambda
add_content_slide(
    "03 › Lambda Function",
    [
        "Lambda — bir qatorli anonim funksiya",
        "  Sintaksis:  lambda parametr: ifoda",
        "Oddiy funksiya vs Lambda:",
        "  def calculate(x, y): return x * y",
        "  lambda x, y: x * y   # bir xil natija",
        "Asosiy ishlatilish: sort key sifatida",
        "  people.sort(key=lambda person: person[1])",
        "  # person[1] = yosh, yoshga ko'ra saralaydi",
    ],
    code='# Oddiy funksiya\ndef calculate(x, y): return x * y\nprint(calculate(5, 10))  # => 50\n\n# people list - (ism, yosh) tuplelar\npeople = [\n    ("Robert", 20),\n    ("Steve", 19),\n    ("Joseph", 25),\n    ("Michael", 30),\n]\n\n# Yoshga ko\'ra saralash\npeople.sort(key=lambda person: person[1])\nprint(people)\n# [("Steve",19),("Robert",20),...]',
    accent_color=ACCENT_YELLOW,
    code_height=3.2
)

# Slide 10 - Section 4
add_section_slide("04", "Enumerate, Map & Filter",
                  "Funksional dasturlash usullari:\nindex+value, transform, filter", ACCENT_PURPLE)

# Slide 11 - Enumerate
add_content_slide(
    "04 › Enumerate — Index va Value birga",
    [
        "enumerate() — index va value ni birga beradi",
        "  Oddiy for: faqat value ni beradi",
        "  enumerate bilan: (index, value) tuple beradi",
        "Dictionaries bilan o'xshashlik:",
        "  car_obj.items() — (key, value) tuplelar",
        "  for (key, value) in car_obj.items():",
        "Ikkalasi ham tuple unpacking ishlatadi",
    ],
    code='animals = ["dog", "cat", "fish"]\n\n# Oddiy for - faqat value\nfor a in animals:\n    print(a)  # dog, cat, fish\n\n# enumerate - index + value\nfor (index, value) in enumerate(animals):\n    print(f"{index}: {value}")\n# 0: dog\n# 1: cat\n# 2: fish\n\n# Dictionary items\ncar = dict(brand="BMW", year=2026)\nfor (key, value) in car.items():\n    print(f"{key}: {value}")',
    accent_color=ACCENT_PURPLE,
    code_height=3.2
)

# Slide 12 - Map & Filter
add_content_slide(
    "04 › Map va Filter",
    [
        "map() — har bir elementga funksiya qo'llaydi",
        "  map(lambda car: car[0], cars)",
        "  # cars listidan faqat nomlarni oladi",
        "filter() — shartga mos elementlarni saqlaydi",
        "  filter(lambda car: car[1] > 80, cars)",
        "  # narxi 80+ bo'lgan mashinalar",
        "MUHIM: map/filter object qaytaradi, list() kerak!",
        "  new_cars = list(map(...))  # ko'rish uchun",
    ],
    code='cars = [\n    ("Ferari", 70), ("Tayota", 87),\n    ("Audi", 116), ("BMW", 109),\n    ("Pagani", 33)\n]\n\n# MAP - nomlarni olish\nnatija_map = map(lambda car: car[0], cars)\nnew_cars = list(natija_map)\nprint(new_cars)\n# ["Ferari","Tayota","Audi","BMW","Pagani"]\n\n# FILTER - narxi > 80\nnatija_filter = filter(\n    lambda car: car[1] > 80, cars)\nprint(list(natija_filter))\n# [("Tayota",87),("Audi",116),("BMW",109)]',
    accent_color=ACCENT_PURPLE,
    code_height=3.8
)

# Slide 13 - Summary
slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(slide_layout)
set_bg(slide)
add_rect(slide, 0, 0, 13.33, 0.08, ACCENT_BLUE)
add_rect(slide, 0, 7.42, 13.33, 0.08, ACCENT_BLUE)
add_rect(slide, 0, 0.08, 13.33, 0.9, RGBColor(22, 26, 48))
add_text_box(slide, "XULOSA — Python Lists", 0.3, 0.1, 12.5, 0.85,
             font_size=26, bold=True, color=WHITE, font_name="Arial")

summary = [
    ("01", "Working with Lists", "Literal [] va Constructor list() — 2 xil yaratish usuli", ACCENT_BLUE),
    ("02", "List Methods", "append/insert/pop/remove (mutable) vs sorted/index (immutable)", ACCENT_GREEN),
    ("03", "Lambda Function", "lambda x: ifoda — anonim funksiya, sort key sifatida", ACCENT_YELLOW),
    ("04", "Enumerate, Map, Filter", "enumerate(index+val) · map(transform) · filter(shart)", ACCENT_PURPLE),
]

for i, (num, title, desc, col) in enumerate(summary):
    y = 1.2 + i * 1.5
    add_rect(slide, 0.3, y, 0.6, 1.2, col)
    add_text_box(slide, num, 0.3, y + 0.25, 0.6, 0.7, font_size=18, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER, font_name="Arial")
    add_rect(slide, 0.95, y, 11.8, 1.2, RGBColor(25, 30, 55))
    add_text_box(slide, title, 1.1, y + 0.1, 11.5, 0.5, font_size=18, bold=True,
                 color=col, font_name="Arial")
    add_text_box(slide, desc, 1.1, y + 0.6, 11.5, 0.5, font_size=14,
                 color=LIGHT_GRAY, font_name="Arial")


output_path = "/Users/ilyosbek96/Desktop/PRACTIC/Python_Lists_Masterclass.pptx"
prs.save(output_path)
print(f"Saved: {output_path}")
